"""اختبارات استخراج النص من الملفات المرفوعة.

الملفات تُبنى داخل الاختبار بالبايت: لا مرفقات ثابتة في المستودع ولا اعتماد
على قاعدة بيانات. ملف PDF يُركَّب يدويًا بأصغر بنية صالحة، وملف DOCX يُنشأ
عبر python-docx.
"""

import io

import pytest

from app.core.config import settings
from app.services.text_extraction import (
    SUPPORTED_EXTENSIONS,
    CorruptFileError,
    EmptyDocumentError,
    FileExtractionError,
    FileTooLargeError,
    UnsupportedFileTypeError,
    extract_text,
    validate_upload,
)

PDF_TEXT = "Official letter draft for the ministry."


def build_pdf(text: str = PDF_TEXT) -> bytes:
    """يبني أصغر ملف PDF صالح يحتوي سطر نص واحد، بجدول xref صحيح."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]"
        b"/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj".encode() + body + b"endobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer<</Size {len(objects) + 1}/Root 1 0 R>>\n".encode()
    out += f"startxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def build_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        created = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for cell_index, value in enumerate(row):
                created.cell(row_index, cell_index).text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 1) رفض الامتدادات غير المدعومة
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "filename", ["تقرير.exe", "بيانات.json", "صورة.png", "أرشيف.zip"]
)
def test_unsupported_extension_is_rejected_with_arabic_message(filename):
    with pytest.raises(UnsupportedFileTypeError) as exc:
        extract_text(filename, b"data")

    message = str(exc.value)
    # الرسالة تذكر الصيغ المقبولة كلها.
    for extension in SUPPORTED_EXTENSIONS:
        assert extension in message


@pytest.mark.parametrize(
    ("filename", "hint"),
    [
        ("خطاب.doc", ".docx"),
        ("جدول.xlsx", "الجداول"),
        ("عرض.pptx", "العروض"),
        ("بيانات.csv", "CSV"),
    ],
)
def test_known_alternatives_get_a_helpful_hint(filename, hint):
    with pytest.raises(UnsupportedFileTypeError) as exc:
        extract_text(filename, b"data")
    assert hint in str(exc.value)


def test_file_without_extension_is_rejected():
    with pytest.raises(UnsupportedFileTypeError) as exc:
        extract_text("مستند_بلا_امتداد", b"data")
    assert "بلا امتداد" in str(exc.value)


def test_extension_check_is_case_insensitive():
    assert extract_text("تقرير.TXT", "نص عربي".encode("utf-8")) == "نص عربي"


def test_rejection_happens_before_reading_the_file():
    """الامتداد يُرفض دون محاولة تفسير المحتوى، مهما كان."""
    with pytest.raises(UnsupportedFileTypeError):
        extract_text("خبيث.exe", b"\x00\xff" * 100)


# ---------------------------------------------------------------------------
# 2) الحجم
# ---------------------------------------------------------------------------
def test_oversized_file_is_rejected(monkeypatch):
    monkeypatch.setattr(settings, "upload_max_bytes", 1024)

    with pytest.raises(FileTooLargeError) as exc:
        extract_text("كبير.txt", b"x" * 2048)

    message = str(exc.value)
    assert "يتجاوز الحد" in message
    assert "كبير.txt" in message


def test_empty_file_is_rejected():
    with pytest.raises(FileExtractionError) as exc:
        extract_text("فارغ.txt", b"")
    assert "فارغ" in str(exc.value)


def test_validate_upload_accepts_a_valid_file():
    assert validate_upload("تقرير.pdf", 5000) == ".pdf"


# ---------------------------------------------------------------------------
# 3) TXT
# ---------------------------------------------------------------------------
def test_txt_arabic_utf8_is_extracted():
    text = "خطاب رسمي\n\nالفقرة الثانية من الخطاب."
    assert extract_text("خطاب.txt", text.encode("utf-8")) == text


def test_txt_with_bom_is_handled():
    assert extract_text("خطاب.txt", "نص".encode("utf-8-sig")) == "نص"


def test_txt_windows_arabic_encoding_is_handled():
    """ملفات ويندوز العربية القديمة بترميز cp1256."""
    assert extract_text("قديم.txt", "تعميم".encode("cp1256")) == "تعميم"


def test_txt_with_only_whitespace_is_rejected():
    with pytest.raises(EmptyDocumentError):
        extract_text("فارغ.txt", b"   \n\n   ")


# ---------------------------------------------------------------------------
# 4) PDF
# ---------------------------------------------------------------------------
def test_pdf_text_is_extracted():
    extracted = extract_text("تقرير.pdf", build_pdf())
    assert "Official letter draft" in extracted


def test_pdf_without_a_text_layer_is_rejected():
    """PDF بلا طبقة نصية (صور ممسوحة) يعطي رسالة واضحة لا انهيارًا."""
    with pytest.raises(EmptyDocumentError) as exc:
        extract_text("ممسوح.pdf", build_pdf(text=" "))
    assert "صورًا ممسوحة" in str(exc.value)


def test_corrupt_pdf_is_rejected_without_crashing():
    with pytest.raises(CorruptFileError) as exc:
        extract_text("تالف.pdf", "%PDF-1.4\nهذا ليس ملفًا صالحًا".encode("utf-8"))
    assert "تالفًا" in str(exc.value)


# ---------------------------------------------------------------------------
# 5) DOCX
# ---------------------------------------------------------------------------
def test_docx_arabic_paragraphs_are_extracted():
    data = build_docx(["بسم الله الرحمن الرحيم", "سعادة المدير المحترم،"])
    extracted = extract_text("خطاب.docx", data)

    assert "بسم الله الرحمن الرحيم" in extracted
    assert "سعادة المدير المحترم،" in extracted


def test_docx_table_content_is_extracted():
    """جداول المحاضر الرسمية جزء من النص ولا تُهمَل."""
    data = build_docx(["محضر اجتماع"], table=[["البند", "القرار"], ["الميزانية", "معتمد"]])
    extracted = extract_text("محضر.docx", data)

    assert "البند" in extracted
    assert "معتمد" in extracted


def test_empty_docx_is_rejected():
    with pytest.raises(EmptyDocumentError):
        extract_text("فارغ.docx", build_docx([]))


def test_corrupt_docx_is_rejected_without_crashing():
    with pytest.raises(CorruptFileError) as exc:
        extract_text("تالف.docx", b"PK\x03\x04" + "ليس ملف وورد".encode("utf-8"))
    assert "تالف" in str(exc.value)
