"""استخراج النص من الملفات المرفوعة: PDF و DOCX و TXT.

استيراد مكتبات القراءة (`pypdf` و `python-docx`) **كسول داخل الدوال**: ملف
نصي بسيط لا يحمّل مكتبة PDF، ويبقى المشروع يقلع حتى لو نقصت مكتبة منها.

كل رفض هنا يرفع خطأً برسالة عربية صالحة للعرض مباشرة على الموظف.
"""

import io
from pathlib import Path

from ..core.config import settings

#: الامتدادات المدعومة في الـMVP.
SUPPORTED_EXTENSIONS: tuple[str, ...] = (".pdf", ".docx", ".txt")

#: امتدادات شائعة نرفضها برسالة تشرح البديل بدل رسالة عامة.
_KNOWN_ALTERNATIVES: dict[str, str] = {
    ".doc": "صيغة Word القديمة غير مدعومة. احفظ الملف بصيغة .docx ثم أعد رفعه.",
    ".rtf": "صيغة RTF غير مدعومة. احفظ الملف بصيغة .docx أو .txt ثم أعد رفعه.",
    ".xlsx": "ملفات الجداول غير مدعومة حاليًا. صدّر المحتوى إلى .pdf أو .txt.",
    ".xls": "ملفات الجداول غير مدعومة حاليًا. صدّر المحتوى إلى .pdf أو .txt.",
    ".pptx": "ملفات العروض غير مدعومة حاليًا. صدّر العرض إلى .pdf ثم ارفعه.",
    ".csv": "ملفات CSV غير مدعومة حاليًا. حوّلها إلى .txt ثم أعد رفعها.",
}

#: ترتيب محاولات فك ترميز الملفات النصية. cp1256 شائع في ملفات ويندوز العربية.
_TEXT_ENCODINGS: tuple[str, ...] = ("utf-8-sig", "utf-8", "cp1256", "cp1252")


class FileExtractionError(Exception):
    """خطأ في معالجة ملف مرفوع، برسالة عربية صالحة للعرض."""


class UnsupportedFileTypeError(FileExtractionError):
    """امتداد غير مدعوم."""


class FileTooLargeError(FileExtractionError):
    """الملف تجاوز الحد الأقصى المسموح."""


class EmptyDocumentError(FileExtractionError):
    """الملف سليم لكن لا نص فيه (صور ممسوحة ضوئيًا مثلًا)."""


class CorruptFileError(FileExtractionError):
    """تعذّر قراءة الملف: تالف أو محمي بكلمة مرور."""


def _human_size(num_bytes: int) -> str:
    megabytes = num_bytes / (1024 * 1024)
    if megabytes >= 1:
        return f"{megabytes:.1f} ميجابايت"
    return f"{num_bytes / 1024:.0f} كيلوبايت"


def file_extension(filename: str) -> str:
    """يعيد امتداد الملف بحروف صغيرة، أو سلسلة فارغة إن لم يكن له امتداد."""
    return Path(filename).suffix.lower()


def validate_upload(filename: str, size_bytes: int) -> str:
    """يتحقق من الامتداد والحجم قبل قراءة الملف، ويعيد الامتداد المعتمد.

    يُستدعى من مسار الرفع قبل تحميل الملف كاملًا في الذاكرة.

    Raises:
        UnsupportedFileTypeError: امتداد غير مدعوم أو بلا امتداد.
        FileTooLargeError: الحجم يتجاوز UPLOAD_MAX_BYTES.
    """
    extension = file_extension(filename)
    supported = "، ".join(SUPPORTED_EXTENSIONS)

    if not extension:
        raise UnsupportedFileTypeError(
            f"الملف «{filename}» بلا امتداد، فتعذّر تحديد نوعه. "
            f"الصيغ المدعومة: {supported}."
        )
    if extension not in SUPPORTED_EXTENSIONS:
        hint = _KNOWN_ALTERNATIVES.get(extension)
        detail = f" {hint}" if hint else ""
        raise UnsupportedFileTypeError(
            f"صيغة الملف «{extension}» غير مدعومة. "
            f"الصيغ المدعومة: {supported}.{detail}"
        )

    if size_bytes <= 0:
        raise FileExtractionError(f"الملف «{filename}» فارغ ولا يحتوي أي بيانات.")

    limit = settings.upload_max_bytes
    if size_bytes > limit:
        raise FileTooLargeError(
            f"حجم الملف «{filename}» ({_human_size(size_bytes)}) يتجاوز الحد "
            f"المسموح ({_human_size(limit)}). قسّم الملف أو اضغطه ثم أعد رفعه."
        )
    return extension


# ---------------------------------------------------------------------------
# القرّاء — استيراد المكتبات كسول داخل كل دالة
# ---------------------------------------------------------------------------
def _read_txt(data: bytes, filename: str) -> str:
    for encoding in _TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise CorruptFileError(
        f"تعذّرت قراءة ترميز الملف النصي «{filename}». "
        "احفظه بترميز UTF-8 ثم أعد رفعه."
    )


def _read_pdf(data: bytes, filename: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise FileExtractionError(
            "مكتبة قراءة PDF غير مثبّتة على هذا السيرفر. "
            "ثبّتها عبر: pip install -r requirements.txt"
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise CorruptFileError(
                f"ملف PDF «{filename}» محمي بكلمة مرور. أزل الحماية ثم أعد رفعه."
            )
        pages = [(page.extract_text() or "") for page in reader.pages]
    except CorruptFileError:
        raise
    except Exception as exc:
        raise CorruptFileError(
            f"تعذّرت قراءة ملف PDF «{filename}». قد يكون تالفًا أو غير مكتمل."
        ) from exc

    return "\n\n".join(page.strip() for page in pages if page.strip())


def _read_docx(data: bytes, filename: str) -> str:
    try:
        import docx
    except ImportError as exc:
        raise FileExtractionError(
            "مكتبة قراءة Word غير مثبّتة على هذا السيرفر. "
            "ثبّتها عبر: pip install -r requirements.txt"
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise CorruptFileError(
            f"تعذّرت قراءة ملف Word «{filename}». تأكد أنه بصيغة .docx وغير تالف."
        ) from exc

    parts = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    # نص الجداول جزء أصيل من الخطابات والمحاضر الرسمية، فلا يُهمَل.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n\n".join(parts)


_READERS = {
    ".txt": _read_txt,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
}


def extract_text(filename: str, data: bytes) -> str:
    """يستخرج نص الملف المرفوع.

    Args:
        filename: اسم الملف الأصلي (يُستخدم لتحديد النوع وللرسائل).
        data: محتوى الملف بالبايت.

    Returns:
        النص المستخرج بعد التنظيف.

    Raises:
        UnsupportedFileTypeError: امتداد غير مدعوم.
        FileTooLargeError: تجاوز الحد الأقصى للحجم.
        EmptyDocumentError: الملف لا يحتوي نصًا قابلًا للاستخراج.
        CorruptFileError: الملف تالف أو محمي.
    """
    extension = validate_upload(filename, len(data))
    text = _READERS[extension](data, filename)

    if not text.strip():
        raise EmptyDocumentError(
            f"لم يُعثر على نص قابل للاستخراج في «{filename}». "
            "قد يكون الملف صورًا ممسوحة ضوئيًا بلا طبقة نصية."
        )
    return text
